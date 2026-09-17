# -*- coding: utf-8 -*-
"""
将 03-拓展功能.md 转换为排版精美的 Word 文档。
"""
import os
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement

# ====== 品牌色 ======
BRAND_BLUE = RGBColor(0x4A, 0x6B, 0xFF)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
MID_GRAY = RGBColor(0x66, 0x66, 0x66)
CODE_BG = "F5F5F5"
TABLE_HEADER_BG = "4A6BFF"
QUOTE_LINE_COLOR = "4A6BFF"

# ====== 字体 ======
FONT_HEADING = "微软雅黑"
FONT_BODY = "微软雅黑"
FONT_CODE = "Consolas"

SRC = r"E:\桌面\鸿蒙报告\03-拓展功能.md"
DST = r"E:\桌面\鸿蒙报告\03-拓展功能.docx"


# ====== 辅助函数 ======
def set_cell_shading(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tc_pr.append(shd)


def set_cell_borders(cell, color="BFBFBF", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{edge}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), size)
        b.set(qn('w:color'), color)
        tc_borders.append(b)
    tc_pr.append(tc_borders)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    """单元格内边距，单位 dxa (1/20 pt)"""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for edge, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        m = OxmlElement(f'w:{edge}')
        m.set(qn('w:w'), str(val))
        m.set(qn('w:type'), 'dxa')
        tc_mar.append(m)
    tc_pr.append(tc_mar)


def set_paragraph_shading(paragraph, color_hex):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    p_pr.append(shd)


def set_paragraph_left_border(paragraph, color_hex="4A6BFF", size="24"):
    """左侧竖线"""
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), size)
    left.set(qn('w:space'), '8')
    left.set(qn('w:color'), color_hex)
    p_bdr.append(left)
    p_pr.append(p_bdr)


def set_paragraph_bottom_border(paragraph, color_hex="4A6BFF", size="12"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size)
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color_hex)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def set_run_font(run, name=FONT_BODY, size=11, bold=False, italic=False, color=None):
    run.font.name = name
    # 中文字体设置
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:ascii'), name)
    rfonts.set(qn('w:hAnsi'), name)
    rfonts.set(qn('w:eastAsia'), name)
    rfonts.set(qn('w:cs'), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_page_break(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_break()
    # 使用 w:br type="page"
    br = run._element.find(qn('w:br'))
    if br is None:
        br = OxmlElement('w:br')
        run._element.append(br)
    br.set(qn('w:type'), 'page')


def add_horizontal_rule(doc, color_hex="4A6BFF", size="12"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    set_paragraph_bottom_border(p, color_hex=color_hex, size=size)
    return p


# ====== 内容渲染 ======
def add_cover(doc):
    # 顶部留白
    for _ in range(6):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)

    # 主标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("NeuShare")
    set_run_font(run, name=FONT_HEADING, size=48, bold=True, color=BRAND_BLUE)
    p.paragraph_format.space_after = Pt(12)

    # 副标题1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("拓展功能文档")
    set_run_font(run, name=FONT_HEADING, size=32, bold=True, color=BRAND_BLUE)
    p.paragraph_format.space_after = Pt(24)

    # 品牌色分隔线
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_bottom_border(p, color_hex="4A6BFF", size="18")
    p.paragraph_format.space_after = Pt(24)

    # 副标题2
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("东北大学校园学习资料分享平台")
    set_run_font(run, name=FONT_HEADING, size=18, bold=False, color=DARK_GRAY)
    p.paragraph_format.space_after = Pt(12)

    # 描述
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("HarmonyOS NEXT · Vue 3 · Spring Boot 3")
    set_run_font(run, name=FONT_BODY, size=12, color=MID_GRAY)
    p.paragraph_format.space_after = Pt(60)

    # 日期
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("2026 年 6 月")
    set_run_font(run, name=FONT_HEADING, size=14, bold=True, color=BRAND_BLUE)

    # 分页
    add_page_break(doc)


def add_heading1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(text)
    set_run_font(run, name=FONT_HEADING, size=16, bold=True, color=BRAND_BLUE)
    # 标题下方加细线
    set_paragraph_bottom_border(p, color_hex="4A6BFF", size="8")


def add_heading2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(text)
    set_run_font(run, name=FONT_HEADING, size=14, bold=True, color=DARK_GRAY)


def add_heading3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(text)
    set_run_font(run, name=FONT_HEADING, size=12, bold=True, color=DARK_GRAY)


def add_body(doc, text):
    """处理正文，支持 **加粗** 与 `代码` 内联"""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.first_line_indent = Pt(0)
    add_inline_runs(p, text, base_size=11)
    return p


def add_inline_runs(p, text, base_size=11, base_color=None):
    """解析 **bold** 与 `code` 内联标记"""
    pattern = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`)')
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            run = p.add_run(text[pos:m.start()])
            set_run_font(run, name=FONT_BODY, size=base_size, color=base_color)
        token = m.group(0)
        if token.startswith('**'):
            run = p.add_run(token[2:-2])
            set_run_font(run, name=FONT_BODY, size=base_size, bold=True, color=base_color)
        elif token.startswith('`'):
            run = p.add_run(token[1:-1])
            set_run_font(run, name=FONT_CODE, size=base_size - 1, color=RGBColor(0xC7, 0x25, 0x4E))
        pos = m.end()
    if pos < len(text):
        run = p.add_run(text[pos:])
        set_run_font(run, name=FONT_BODY, size=base_size, color=base_color)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.74 + level * 0.5)
    add_inline_runs(p, text, base_size=11)


def add_numbered(doc, text, level=0):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.74 + level * 0.5)
    add_inline_runs(p, text, base_size=11)


def add_quote(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    set_paragraph_left_border(p, color_hex=QUOTE_LINE_COLOR, size="24")
    # 引用内可能含 **bold**
    add_inline_runs(p, text, base_size=11, base_color=MID_GRAY)
    for run in p.runs:
        run.italic = True


def add_code_block(doc, code_text):
    """代码块：Consolas 10pt，浅灰背景"""
    lines = code_text.split('\n')
    # 用单段落 + 多 run + 换行，便于整段加背景
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.right_indent = Cm(0.3)
    set_paragraph_shading(p, CODE_BG)
    # 边框
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement('w:pBdr')
    for edge in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{edge}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '4')
        b.set(qn('w:color'), 'DDDDDD')
        p_bdr.append(b)
    p_pr.append(p_bdr)

    for i, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, name=FONT_CODE, size=10, color=RGBColor(0x1F, 0x23, 0x2E))
        if i < len(lines) - 1:
            run.add_break()


def add_table(doc, header, rows):
    """渲染表格：表头蓝底白字，网格边框"""
    n_cols = len(header)
    table = doc.add_table(rows=1, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # 表头
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(header):
        cell = hdr_cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_run_font(run, name=FONT_HEADING, size=10.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        set_cell_shading(cell, TABLE_HEADER_BG)
        set_cell_borders(cell, color="4A6BFF", size="6")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # 数据行
    for r_idx, row_data in enumerate(rows):
        row = table.add_row()
        for i, val in enumerate(row_data):
            cell = row.cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
            # 内联解析
            add_inline_runs(p, val, base_size=10)
            # 隔行底色
            if r_idx % 2 == 1:
                set_cell_shading(cell, "F2F5FF")
            set_cell_borders(cell, color="BFBFBF", size="4")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # 表格后空行
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ====== Markdown 解析 ======
def parse_markdown(md_text):
    """将 markdown 解析为指令列表"""
    lines = md_text.split('\n')
    i = 0
    blocks = []
    n = len(lines)

    while i < n:
        line = lines[i]
        stripped = line.rstrip()

        # 跳过空行
        if not stripped.strip():
            i += 1
            continue

        # 水平分隔线
        if stripped.strip() in ('---', '***', '___'):
            blocks.append(('hr', None))
            i += 1
            continue

        # 代码块
        if stripped.startswith('```'):
            code_lines = []
            i += 1
            while i < n and not lines[i].startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1  # 跳过结束 ```
            blocks.append(('code', '\n'.join(code_lines)))
            continue

        # 标题
        if stripped.startswith('#'):
            m = re.match(r'^(#{1,6})\s+(.*)$', stripped)
            if m:
                level = len(m.group(1))
                text = m.group(2).strip()
                blocks.append(('h', (level, text)))
                i += 1
                continue

        # 引用块（可能多行）
        if stripped.startswith('>'):
            quote_lines = []
            while i < n and lines[i].lstrip().startswith('>'):
                quote_lines.append(re.sub(r'^>\s?', '', lines[i].lstrip()))
                i += 1
            blocks.append(('quote', '\n'.join(quote_lines)))
            continue

        # 表格
        if '|' in stripped and i + 1 < n and re.match(r'^\s*\|?[\s\-:|]+\|?\s*$', lines[i + 1]):
            # 解析表格
            header_line = stripped
            i += 1  # 跳过分隔行
            i += 1
            rows = []
            while i < n and '|' in lines[i] and lines[i].strip():
                rows.append(lines[i].strip())
                i += 1
            blocks.append(('table', (header_line, rows)))
            continue

        # 有序列表
        if re.match(r'^\d+\.\s+', stripped):
            list_items = []
            while i < n and re.match(r'^\d+\.\s+', lines[i].lstrip()):
                item = re.sub(r'^\d+\.\s+', '', lines[i].lstrip())
                list_items.append(item)
                i += 1
            blocks.append(('ol', list_items))
            continue

        # 无序列表
        if re.match(r'^[-*+]\s+', stripped):
            list_items = []
            while i < n and re.match(r'^[-*+]\s+', lines[i].lstrip()):
                item = re.sub(r'^[-*+]\s+', '', lines[i].lstrip())
                list_items.append(item)
                i += 1
            blocks.append(('ul', list_items))
            continue

        # 普通段落（连续非空行合并）
        para_lines = [stripped]
        i += 1
        while i < n:
            nxt = lines[i].rstrip()
            if not nxt.strip():
                break
            if (nxt.startswith('#') or nxt.startswith('>') or nxt.startswith('```')
                    or nxt in ('---', '***', '___')
                    or re.match(r'^\d+\.\s+', nxt)
                    or re.match(r'^[-*+]\s+', nxt)
                    or ('|' in nxt and i + 1 < n and re.match(r'^\s*\|?[\s\-:|]+\|?\s*$', lines[i + 1]))):
                break
            para_lines.append(nxt)
            i += 1
        blocks.append(('p', ' '.join(para_lines)))

    return blocks


def parse_table_row(line):
    """解析表格行 -> 单元格列表"""
    line = line.strip()
    if line.startswith('|'):
        line = line[1:]
    if line.endswith('|'):
        line = line[:-1]
    return [c.strip() for c in line.split('|')]


# ====== 页脚 ======
def setup_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 文档名 + 日期
    run = p.add_run("NeuShare 拓展功能文档  ·  2026 年 6 月")
    set_run_font(run, name=FONT_BODY, size=9, color=MID_GRAY)


# ====== 主流程 ======
def main():
    with open(SRC, 'r', encoding='utf-8') as f:
        md_text = f.read()

    doc = Document()

    # 页面设置
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # 默认正文样式
    style = doc.styles['Normal']
    style.font.name = FONT_BODY
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:ascii'), FONT_BODY)
    rfonts.set(qn('w:hAnsi'), FONT_BODY)
    rfonts.set(qn('w:eastAsia'), FONT_BODY)

    # 封面
    add_cover(doc)

    # 页脚
    setup_footer(doc)

    # 解析 markdown
    blocks = parse_markdown(md_text)

    # 跳过第一个 H1（封面已用），从目录或正文开始
    # 找到第一个非 H1 块开始渲染
    start_idx = 0
    for idx, (kind, _) in enumerate(blocks):
        if kind == 'h':
            # 跳过顶级 H1
            continue
        start_idx = idx
        break

    # 实际上我们想保留目录与所有内容，只是不重复 H1 标题
    # 重新策略：渲染所有块，但 H1 级别用封面替代（跳过首个 H1）
    h1_skipped = False

    for kind, data in blocks:
        if kind == 'h':
            level, text = data
            if level == 1:
                if not h1_skipped:
                    h1_skipped = True
                    continue  # 封面已展示
                add_heading1(doc, text)
            elif level == 2:
                add_heading2(doc, text)
            elif level == 3:
                add_heading3(doc, text)
            else:
                add_heading3(doc, text)
        elif kind == 'p':
            add_body(doc, data)
        elif kind == 'ul':
            for item in data:
                add_bullet(doc, item)
        elif kind == 'ol':
            for item in data:
                add_numbered(doc, item)
        elif kind == 'quote':
            # 引用可能多行，每行单独渲染
            for q_line in data.split('\n'):
                if q_line.strip():
                    add_quote(doc, q_line)
        elif kind == 'code':
            add_code_block(doc, data)
        elif kind == 'table':
            header_line, rows_lines = data
            header = parse_table_row(header_line)
            rows = [parse_table_row(r) for r in rows_lines]
            add_table(doc, header, rows)
        elif kind == 'hr':
            add_horizontal_rule(doc, color_hex="4A6BFF", size="12")

    # 保存
    os.makedirs(os.path.dirname(DST), exist_ok=True)
    doc.save(DST)
    print(f"✅ 已生成: {DST}")
    print(f"   文件大小: {os.path.getsize(DST) / 1024:.1f} KB")


if __name__ == '__main__':
    main()
